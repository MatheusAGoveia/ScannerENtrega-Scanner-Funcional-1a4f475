from __future__ import annotations

import csv
import io
import re
import zipfile
from pathlib import PurePosixPath

from openpyxl import load_workbook
from pypdf import PdfReader


class ImportValidationError(ValueError):
    pass


SUPPORTED_EXTENSIONS = {".txt", ".csv", ".xlsx", ".pdf", ".zip"}
TOKEN_SPLIT = re.compile(r"[\s,;|]+")


def _text_tokens(text: str) -> list[str]:
    return [token.strip() for token in TOKEN_SPLIT.split(text) if token.strip()]


def _parse_text(content: bytes) -> list[str]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    if re.search(r"(?im)^\s*(=|\+|@|cmd\||powershell|javascript:)", text):
        raise ImportValidationError("Conteudo executavel ou formula nao e permitido.")
    return _text_tokens(text)


def _parse_csv(content: bytes) -> list[str]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    rows: list[str] = []
    for row in csv.reader(io.StringIO(text)):
        for value in row:
            rows.extend(_text_tokens(value))
    return rows


def _parse_xlsx(content: bytes) -> list[str]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ImportValidationError("Planilha XLSX invalida ou corrompida.") from exc
    values: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        raise ImportValidationError("Planilhas com formulas nao sao permitidas.")
                    if cell.value is not None:
                        values.extend(_text_tokens(str(cell.value)))
    finally:
        workbook.close()
    return values


def _parse_pdf(content: bytes) -> list[str]:
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ImportValidationError("PDF invalido ou corrompido.") from exc
    if not text.strip():
        raise ImportValidationError(
            "O PDF nao possui camada de texto. OCR nao e executado no servidor."
        )
    return _text_tokens(text)


def parse_import_file(filename: str, content: bytes, *, max_bytes: int) -> list[str]:
    if not content:
        raise ImportValidationError("O arquivo esta vazio.")
    if len(content) > max_bytes:
        raise ImportValidationError("O arquivo excede o limite configurado.")
    if content.startswith((b"MZ", b"\x7fELF")):
        raise ImportValidationError("Arquivos executaveis nao sao permitidos.")

    extension = PurePosixPath(filename.lower()).suffix
    if extension not in SUPPORTED_EXTENSIONS:
        raise ImportValidationError("Formato nao suportado. Use TXT, CSV, XLSX, PDF ou ZIP.")
    if extension == ".txt":
        return _parse_text(content)
    if extension == ".csv":
        return _parse_csv(content)
    if extension == ".xlsx":
        return _parse_xlsx(content)
    if extension == ".pdf":
        return _parse_pdf(content)

    values: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            safe_entries = [entry for entry in archive.infolist() if not entry.is_dir()]
            if len(safe_entries) > 50:
                raise ImportValidationError("O ZIP excede o limite de 50 arquivos.")
            total_uncompressed = sum(entry.file_size for entry in safe_entries)
            if total_uncompressed > max_bytes * 4:
                raise ImportValidationError(
                    "O conteudo descompactado excede o limite de seguranca."
                )
            for entry in safe_entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ImportValidationError("Caminho inseguro detectado dentro do ZIP.")
                if path.suffix.lower() not in SUPPORTED_EXTENSIONS - {".zip"}:
                    continue
                values.extend(
                    parse_import_file(path.name, archive.read(entry), max_bytes=max_bytes)
                )
    except zipfile.BadZipFile as exc:
        raise ImportValidationError("Arquivo ZIP invalido ou corrompido.") from exc
    return values
